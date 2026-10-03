"""M15 Audit & compliance: immutability (ORM / queryset / API / raw SQL), site derivation, tenant + site scoped
viewer, filters / categories / entity history, exports (contents reconciled with the database), export authorization
and the audit-of-the-audit."""
import csv
import io
import uuid

import pytest
from django.db import IntegrityError, connection, transaction
from django.test import Client
from rest_framework.test import APIClient

from apps.audit import exports
from apps.audit import services as audit
from apps.audit.models import AuditLog, ImmutableError
from apps.incidents import services as incidents
from tests.phase3_support import new_wo

pytestmark = pytest.mark.django_db


def api(user, org):
    c = APIClient()
    c.force_authenticate(user=user)
    c.credentials(HTTP_X_ORGANIZATION=org.slug)
    return c


def web(user):
    c = Client()
    c.force_login(user)
    return c


@pytest.fixture
def a_(p3, org_b, make_member, make_site, make_asset):
    """Real activity in two sites of Alpha and one site of Beta (everything through the services)."""
    wo1 = new_wo(p3, title="Pump service")
    sr = incidents.create_request(p3["org"], asset=p3["asset"], reporter=p3["tech"], title="Leak", severity="HIGH",
                                  actor=p3["tech"].user)
    incidents.transition(sr, action="triage", actor=p3["ops"].user)
    incidents.transition(sr, action="approve", actor=p3["ops"].user)
    wo2 = new_wo(p3, asset=p3["asset2"], title="Fan service")
    b_ops = make_member(org_b, "ops@beta.test", "operations_manager")
    b_reader = make_member(org_b, "auditor@beta.test", "auditor")
    b_asset = make_asset(org_b, make_site(org_b, "B4"), "B-FAN")
    from apps.workorders import services as wos

    b_wo = wos.create_work_order(org_b, asset=b_asset, actor=b_ops.user, title="Beta job", work_type="CORRECTIVE")
    return {**p3, "wo1": wo1, "wo2": wo2, "sr": sr, "b_ops": b_ops, "b_reader": b_reader, "b_wo": b_wo,
            "org_b": org_b}


# --- immutability ------------------------------------------------------------------------------------------------


def test_audit_rows_cannot_be_changed_or_deleted_by_any_path(a_):
    row = AuditLog.objects.for_organization(a_["org"]).first()
    row.action = "tampered"
    with pytest.raises(ImmutableError):
        row.save()
    with pytest.raises(ImmutableError):
        row.delete()
    with pytest.raises(ImmutableError):
        AuditLog.objects.filter(pk=row.pk).update(action="tampered")
    with pytest.raises(ImmutableError):
        AuditLog.objects.filter(pk=row.pk).delete()
    with pytest.raises(ImmutableError):
        AuditLog.objects.all().delete()
    for sql in ("UPDATE audit_auditlog SET action = 'tampered' WHERE id = %s",
                "DELETE FROM audit_auditlog WHERE id = %s"):
        with pytest.raises(IntegrityError), transaction.atomic(), connection.cursor() as cur:
            cur.execute(sql, [str(row.pk)])
    with pytest.raises(IntegrityError), transaction.atomic(), connection.cursor() as cur:
        cur.execute("UPDATE audit_auditlog SET action = 'x'")  # even a mass update
    fresh = AuditLog.objects.get(pk=row.pk)
    assert fresh.action != "tampered"


def test_api_offers_no_write_routes_for_anyone(a_):
    row = AuditLog.objects.for_organization(a_["org"]).first()
    owner_like = api(a_["reader"].user, a_["org"])
    for method in ("post", "put", "patch", "delete"):
        for url in ("/api/v1/audit-logs/", f"/api/v1/audit-logs/{row.pk}/"):
            r = getattr(owner_like, method)(url, {"action": "x"}, format="json")
            assert r.status_code in (403, 404, 405), (method, url, r.status_code)
    assert AuditLog.objects.get(pk=row.pk).action == row.action


# --- site derivation ---------------------------------------------------------------------------------------------


def test_rows_record_the_site_of_the_audited_record(a_):
    p = a_
    wo_row = AuditLog.objects.get(action="work_order.created", target_id=str(p["wo1"].pk))
    assert wo_row.site_id == p["site"].pk
    assert AuditLog.objects.get(action="work_order.created", target_id=str(p["wo2"].pk)).site_id == p["site2"].pk
    assert AuditLog.objects.filter(action="incident.created", target_id=str(p["sr"].pk)).get().site_id == p["site"].pk
    org_row = AuditLog.objects.filter(organization=p["org"], action="organization.created").first()
    assert org_row is None or org_row.site_id is None


# --- viewer scope, tenant isolation ---------------------------------------------------------------------------


def test_viewer_is_tenant_and_site_scoped(a_, make_scoped_member):
    p = a_
    ids_alpha = {str(i) for i in AuditLog.objects.filter(organization=p["org"]).values_list("pk", flat=True)}
    full = api(p["reader"].user, p["org"]).get("/api/v1/audit-logs/", {"page_size": 200}).json()
    got = {r["id"] for r in full["results"]}
    assert got <= ids_alpha and full["count"] == len(ids_alpha)
    beta_row = AuditLog.objects.get(action="work_order.created", target_id=str(p["b_wo"].pk))
    assert str(beta_row.pk) not in got
    assert api(p["reader"].user, p["org"]).get(f"/api/v1/audit-logs/{beta_row.pk}/").status_code == 404
    assert web(p["reader"].user).get(f"/app/audit/{beta_row.pk}/").status_code == 404
    scoped = make_scoped_member(p["org"], "aud2@alpha.test", "auditor", [p["site2"]])
    r = api(scoped.user, p["org"]).get("/api/v1/audit-logs/", {"page_size": 200}).json()
    assert r["count"] >= 1
    assert all(row["site_id"] == str(p["site2"].pk) for row in r["results"])  # organization-level rows are hidden too
    site1_row = AuditLog.objects.get(action="work_order.created", target_id=str(p["wo1"].pk))
    assert api(scoped.user, p["org"]).get(f"/api/v1/audit-logs/{site1_row.pk}/").status_code == 404
    assert web(scoped.user).get(f"/app/audit/{site1_row.pk}/").status_code == 404
    assert api(scoped.user, p["org"]).get("/api/v1/audit-logs/", {"site": str(p["site"].pk)}).json()["count"] == 0


def test_roles_without_audit_view_are_refused(a_):
    for who in ("tech", "planner", "sup"):
        assert api(a_[who].user, a_["org"]).get("/api/v1/audit-logs/").status_code == 403, who
        assert web(a_[who].user).get("/app/audit/").status_code == 403, who
    assert web(a_["reader"].user).get("/app/audit/").status_code == 200
    assert APIClient().get("/api/v1/audit-logs/").status_code in (401, 403)


# --- filters, categories, histories -------------------------------------------------------------------------


def test_filters_and_categories(a_):
    p = a_
    c = api(p["reader"].user, p["org"])

    def rows(**params):
        return c.get("/api/v1/audit-logs/", {"page_size": 200, **params}).json()["results"]

    assert {r["action"] for r in rows(action="work_order.")} == {"work_order.created"}
    assert all(r["actor_email"] == "ops@alpha.test" for r in rows(actor="ops@alpha"))
    assert {r["target_id"] for r in rows(entity_type="workorders.workorder", action="work_order.created")} == {
        str(p["wo1"].pk), str(p["wo2"].pk)}
    assert len(rows(entity_id=str(p["sr"].pk))) >= 3
    approvals = rows(category="approvals")
    assert approvals and all(r["metadata"]["action"] in ("triage", "approve", "reject", "confirm", "reopen", "review",
                                                         "start_review") for r in approvals)
    assert {r["metadata"]["action"] for r in approvals} == {"triage", "approve"}
    assert rows(category="closures") == []
    assert {r["action"] for r in rows(category="state")} >= {"incident.status_changed"}
    assert rows(category="inventory") == []
    assert rows(**{"from": "2999-01-01"}) == [] and len(rows(to="2999-01-01")) > 0
    assert c.get("/api/v1/audit-logs/", {"from": "2026-02-31"}).status_code == 400
    assert c.get("/api/v1/audit-logs/", {"from": "2026-05-02", "to": "2026-05-01"}).status_code == 400
    assert c.get("/api/v1/audit-logs/", {"category": "bogus"}).status_code == 400
    assert c.get("/api/v1/audit-logs/", {"site": "nope"}).status_code == 400


def test_security_category_records_failed_scans_and_logins(a_):
    from apps.identification import services as idn

    idn.resolve(a_["tech"].user, "nonexistent-token-12345", a_["tech"])
    sec = api(a_["reader"].user, a_["org"]).get("/api/v1/audit-logs/", {"category": "security"}).json()["results"]
    assert "qr.scan_unknown" in {r["action"] for r in sec}


def test_asset_and_work_order_histories(a_):
    p = a_
    c = api(p["reader"].user, p["org"])
    by_tag = c.get("/api/v1/audit-logs/", {"asset": "P3-PUMP", "page_size": 200}).json()["results"]
    kinds = {r["target_type"] for r in by_tag}
    assert {"workorders.workorder", "incidents.servicerequest"} <= kinds
    assert str(p["wo2"].pk) not in {r["target_id"] for r in by_tag}  # the other asset's order is not in this history
    by_id = c.get("/api/v1/audit-logs/", {"asset": str(p["asset"].pk), "page_size": 200}).json()["count"]
    assert by_id == len(by_tag)
    wo = c.get("/api/v1/audit-logs/", {"work_order": p["wo1"].number, "page_size": 200}).json()["results"]
    assert wo and {r["target_id"] for r in wo if r["target_type"] == "workorders.workorder"} == {str(p["wo1"].pk)}
    assert c.get("/api/v1/audit-logs/", {"asset": "NO-SUCH-TAG"}).json()["count"] == 0
    assert c.get("/api/v1/audit-logs/", {"asset": str(p["b_wo"].asset_id)}).json()["count"] == 0  # foreign asset


# --- exports -------------------------------------------------------------------------------------------------


def _csv_rows(response):
    text = response.content.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def test_csv_export_matches_the_database_and_the_filters(a_):
    p = a_
    r = web(p["reader"].user).get("/app/audit/export/?format=csv&action=work_order.")
    assert r.status_code == 200 and r["Content-Type"].startswith("text/csv")
    assert r["Content-Disposition"].startswith("attachment") and r["Cache-Control"] == "private, no-store"
    rows = _csv_rows(r)
    assert rows[0] == exports.COLUMNS
    expected = AuditLog.objects.filter(organization=p["org"], action__startswith="work_order.")
    assert len(rows) - 1 == expected.count() == 2
    assert {row[4] for row in rows[1:]} == {str(p["wo1"].pk), str(p["wo2"].pk)}
    assert str(p["b_wo"].pk) not in r.content.decode("utf-8-sig")  # no cross-tenant row
    assert {row[6] for row in rows[1:]} == {"A1", "A2"}
    client = web(p["reader"].user)  # logging in writes its own audit row, so count afterwards
    total = AuditLog.objects.filter(organization=p["org"]).count()
    full = _csv_rows(client.get("/app/audit/export/?format=csv"))
    assert len(full) - 1 == total  # the export is audited AFTER it is built


def test_xlsx_and_pdf_exports_are_real_files(a_):
    from openpyxl import load_workbook

    p = a_
    x = web(p["reader"].user).get("/app/audit/export/?format=xlsx&category=approvals")
    assert x.status_code == 200 and "spreadsheetml" in x["Content-Type"]
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb["Audit"]
    data = [[c.value for c in row] for row in ws.iter_rows()]
    assert data[0] == exports.COLUMNS and len(data) - 1 == AuditLog.objects.filter(
        organization=p["org"], action="incident.status_changed", metadata__action__in=["triage", "approve"]).count()
    assert "Organization: Alpha Industries" in [r[0].value for r in wb["About"].iter_rows()]
    pdf = web(p["reader"].user).get("/app/audit/export/?format=pdf&category=approvals")
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-" and pdf["Content-Type"] == "application/pdf"
    assert web(p["reader"].user).get("/app/audit/export/?format=exe").status_code == 400
    assert web(p["reader"].user).get("/app/audit/export/?format=csv&from=zzz").status_code == 400


def test_exports_neutralise_spreadsheet_formulas(a_):
    p = a_
    audit.record("test.formula", actor=p["ops"].user, organization=p["org"],
                 target_repr="=HYPERLINK(\"http://evil.test\",\"x\")")
    rows = _csv_rows(web(p["reader"].user).get("/app/audit/export/?format=csv&action=test.formula"))
    assert any(cell.startswith("'=HYPERLINK") for row in rows for cell in row)
    assert not any(cell.startswith("=") for row in rows for cell in row)


def test_export_authorization_scope_and_self_audit(a_, make_scoped_member, monkeypatch):
    p = a_
    for who in ("ops", "tech", "planner"):  # ops can view (via *.view) but not export
        assert web(p[who].user).get("/app/audit/export/?format=csv").status_code == 403, who
        assert api(p[who].user, p["org"]).get("/api/v1/audit-logs/export/").status_code == 403, who
    assert web(p["ops"].user).get("/app/audit/").status_code == 200
    assert "Export results" not in web(p["ops"].user).get("/app/audit/").content.decode()
    assert "Export results" in web(p["reader"].user).get("/app/audit/").content.decode()
    scoped = make_scoped_member(p["org"], "aud3@alpha.test", "auditor", [p["site2"]])
    rows = _csv_rows(web(scoped.user).get("/app/audit/export/?format=csv"))
    assert rows[1:] and {row[6] for row in rows[1:]} == {"A2"}
    assert str(p["wo1"].pk) not in "".join(",".join(r) for r in rows)
    before = AuditLog.objects.filter(action="audit.exported", organization=p["org"]).count()
    r = api(p["reader"].user, p["org"]).get("/api/v1/audit-logs/export/", {"export_format": "csv", "category": "state"})
    assert r.status_code == 200
    ev = AuditLog.objects.filter(action="audit.exported", organization=p["org"]).order_by("-occurred_at").first()
    assert AuditLog.objects.filter(action="audit.exported", organization=p["org"]).count() == before + 1
    assert ev.metadata["format"] == "csv" and ev.metadata["filters"] == {"category": "state"} and ev.actor == p[
        "reader"].user
    monkeypatch.setattr(exports, "EXPORT_LIMIT", 3)
    t = web(p["reader"].user).get("/app/audit/export/?format=csv")
    assert len(_csv_rows(t)) - 1 == 3
    ev = AuditLog.objects.filter(action="audit.exported").order_by("-occurred_at").first()
    assert ev.metadata["truncated"] is True and ev.metadata["rows"] == 3


# --- pages ----------------------------------------------------------------------------------------------------


def test_pages_render_and_show_validation(a_):
    cl = web(a_["reader"].user)
    for url in ("/app/audit/", "/app/audit/reports/", "/app/audit/?category=approvals", "/app/audit/?asset=P3-PUMP",
                f"/app/audit/?work_order={a_['wo1'].number}", "/app/audit/?q=leak&entity_type=incidents.servicerequest"):
        assert cl.get(url).status_code == 200, url
    row = AuditLog.objects.get(action="incident.created", target_id=str(a_["sr"].pk))
    detail = cl.get(f"/app/audit/{row.pk}/").content.decode()
    assert "incident.created" in detail and "A1" in detail
    bad = cl.get("/app/audit/?from=2026-99-99").content.decode()
    assert "must be a date" in bad
    reports = cl.get("/app/audit/reports/").content.decode()
    for label in ("Security events", "Approvals", "Closures", "Inventory", "SLA events", "Asset history"):
        assert label in reports
    assert cl.get(f"/app/audit/{uuid.uuid4()}/").status_code == 404
